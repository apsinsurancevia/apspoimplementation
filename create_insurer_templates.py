"""
create_insurer_templates.py
----------------------------
Creates 12 insurer statement template Excel files for APS Insurance Broking
Services Pvt Ltd, FY2026-27. Columns are EXACT — names pasted verbatim from
each insurer's original statement file. No cleaning, no renaming.

Save to /mnt/user-data/outputs/

Requirements:
    pip install openpyxl
"""

import os
import sys
import traceback

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Formatting constants
# ---------------------------------------------------------------------------
HEADER_FILL  = PatternFill(fill_type="solid", fgColor="1F3864")
HEADER_FONT  = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
MIN_COL_WIDTH = 15
MAX_COL_WIDTH = 60

OUT_DIR = "/mnt/user-data/outputs"


def col_width(header: str) -> float:
    return min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, len(str(header)) + 4))


def write_sheet(wb: Workbook, sheet_name: str, columns: list,
                blank_rows_above: int = 0) -> None:
    ws = wb.create_sheet(title=sheet_name)

    for _ in range(blank_rows_above):
        ws.append([])

    header_row = blank_rows_above + 1
    ws.append(columns)

    for col_idx, header in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font  = HEADER_FONT
        cell.fill  = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width(
            header if header else "col"
        )

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(columns))}{header_row}"
    )


# ---------------------------------------------------------------------------
# Master configuration — exact columns per insurer per sheet
# ---------------------------------------------------------------------------

INSURERS = [
    {
        "name": "BAJAJ ALLIANZ",
        "file": "BAJAJ_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "irdaPayoutReport",
                "blank_rows_above": 4,
                "columns": [
                    "COMM_PROCESS_DATE","POLICY_REFERENCE","POLICY DATE","MAX COMM RATE",
                    "ACTUAL COMM RATE","TERROR COMM RATE","TP RATE","CUMM COMM RATE",
                    "PREMIUM FOR COMMISSION OD","PREMIUM FOR COMMISSION TP","TERROR PREMIUM",
                    "COMM AMT OD","COMM AMT TP","TERROR COMM AMT","TOTAL COMMISSION",
                    "NET PREMIUM","CONTRACT ID","IMD CODE","PRODUCT","OLD PRODUCT","LOCATION",
                    "PAYMENT DATE","COINS IND","RECOVERY DATE","CHANNEL CODE","CUSTOMER NAME",
                    "PAYOUT ID","CHEQUE NO","COMM DISC RATE","COMM DISC AMT","MOVEMENT",
                    "RID OF POLICY","VEHICLE REGISTRATION YEAR","VEHICLE REG DATE","RID DATE",
                    "POLICY YEAR","RI START DATE",
                ],
            },
        ],
    },
    {
        "name": "GO DIGIT",
        "file": "GODIGIT_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "AAZCA5637P",
                "blank_rows_above": 0,
                "columns": [
                    "policy_number","version_no","endorsement_ind","policy_issue_date","risk_inc_date",
                    "risk_exp_date","office_code","office_name","state","user_id","imd_code","imd_name",
                    "imd_channel","imd_channel_irda","product_code","product_name","product_lob",
                    "vehicle_make","vehicle_model","fuel_type","veh_reg_no","rto_location","veh_age",
                    "veh_type","vehicle_subtype","veh_usage","veh_wheels","veh_seating","veh_gvw",
                    "veh_cc","policy_holder","ph_pincode","od_premium","tp_premium","add_on_premium",
                    "sum_insured","net_premium","issue_source","account_manager","prev_policy_expiry",
                    "curr_ncb","comm_disc_rate","comm_disc_amt","zero_dep_cover","tarrif_premium",
                    "veh_reg_date","technical_premium","cng_flag","SUR_PERS","partner_preferred_discount",
                    "incremental_variation","office_gst_number","ph_gst_number","vehicle_code",
                    "od_premium_before_commerical_discount","od_premium_at_commerical_discount",
                    "carrier_type","net_premium_coll","product_package","insured_person","policy_type",
                    "policy_status","prev_insurer","prev_policy_no","prev_ncb","p'cover",
                    "endorse_reasons","AM_Location","Month","transaction_seq_number","Terrorism",
                    "Pan Number","mail_id","IRDA_AMT","REWARD_AMT","TOTAL_AMT","Master_Policy_No","Gst_Flag",
                ],
            },
        ],
    },
    {
        "name": "ICICI LOMBARD",
        "file": "ICICI_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "","AGENT_CATEGORY_NAME","AGENT_LICENSE_CODE","INTERMEDIARY_CODE","INTERMEDIARY_NAME",
                    "LOCATION_NAME","CUSTOMER_ID","INSURED_CUSTOMER_NAME","PRODUCT_CODE","PRODUCT_NAME",
                    "POLICY_NUMBER","POL_NUM_TXT","POL_PRI_VVERTICAL_DESC","POL_PRI_SSUB_VERTICAL_DESC",
                    "DEAL_CODE","POL_RM_CODE","POL_RM_NAME","APPLICABLE_PREMIUM_AMOUNT",
                    "PREMIUM_FOR_PAYOUTS","POL_COMMISSION_PERCENTAGE","COMMISSION_RATE_USED",
                    "ACTUAL_COMPR_OD_COMM","ACTUAL_COMPR_TP_COMM","ACTUAL_COMMISSION",
                    "POLICY_START_DATE","POLICY_END_DATE","ENDORSEMENT_NUMBER","ENDORSEMENT_TYPE",
                    "LEADER_FOLLOWER","CHILD_INTERMEDIARY_ID","CHILD_INTEREDIARY_NAME",
                    "COMMISSIONABLE_SOURCE","PAYMENT_LOCATION","BRANCH_NAME","BASIC_PREMIUM_AMOUNT",
                    "EARTHQUAKE_PREMIUM_AMOUNT","MOTOR_OD_PREMIUM_AMOUNT","MOTOR_TP_PREMIUM_AMOUNT",
                    "TERRORISM_PREMIUM_AMOUNT","WAR_SRCC_PREMIUM_AMT","PREMIUM_AMOUNT_USED",
                    "BUSINESS_TYPE","IRDA_REPORTING_DATE","GWP","FROM_LOCATION","TO_LOCATION",
                    "POL_VEHICLE_AGE","CUST_TYPE","TRANSACTION_STATUS","YEAR_FLAG","PAYOUT_CYCLE",
                    "PAY_ON","MATCH_TYPE","CALCULATION_TYPE","STD_REMARKS","TRANSACTION_ID",
                    "TOTAL_OD_PREMIUM_RECEIVED","TOTAL_TP_PREMIUM_RECEIVED","TOTAL_PREMIUM_RECEIVED",
                    "DMA_CODE","BRANCH_STATE","OFFLINE_POL_NUM_TXT","BANCS_POL_NUM",
                    "GST_EXCEMPTION_FLAG","OLD_FROM_LOCATION",
                ],
            },
        ],
    },
    {
        "name": "KOTAK",
        "file": "KOTAK_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "BRANCH CODE","BRANCH NAME","ZKGI GST NO","POLICY NO","Endorsement No",
                    "ENDORSEMENT TYPE","Endorsement Reason","BUSINESS TYPE",
                    "Policy/Endorsement Issue Date","POLICY INCEPTION DATE",
                    "Endorsement Effective Date","POLICY EXPIRY DATE","LOB NAME","PRODUCT NAME",
                    "CUSTOMER NAME","IMD CODE","IMD NAME","IMD GSTIN No",
                    "PA TO OWNER DRIVER PREMIUM","ALL OTHER TP COVER PREMIUM","Terrorism Premium",
                    "BASIC PREMIUM/OD PREMIUM","OUR SHARE OF PREMIUM GWP","GST Exempted Yes/No",
                    "TDS RATE ON COMM","PRODUCT TYPE","OD %","TP %","Terr %",
                    "OD / Basic commission","Motor TP commission","Terrisom commission",
                    "Total Commission","TDS Amt","CGST","SGST","IGST","Net Commission",
                ],
            },
        ],
    },
    {
        "name": "NATIONAL INSURANCE",
        "file": "NATIONAL_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "OO Code","UW Year","Month","Department","Sales Channel Type",
                    "Sales Channel Code","Policy Number","Endorsement Number","Bill No",
                    "Issue Date","Effective Date","Expiry Date","Insured Name",
                    "Age of Vehicle","Age of (Insured)","Coinsurance Share","Premium",
                    "Commission","Commission %","Incentive Amount","ACTIVE_FLAG",
                ],
            },
        ],
    },
    {
        "name": "NEW INDIA ASSURANCE",
        "file": "NEWINDIA_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Policy Detail",
                "blank_rows_above": 0,
                "columns": [
                    "Invoice Number","Invoice Date","Office Code","Policy Number",
                    "Endorsement Number","LOB Code","Insured Name","Insured Type",
                    "Policy Month","Policy Inception Date","Policy Expiry Date",
                    "Brokerage Percentage","Premium","Brokerage",
                ],
            },
            {
                "name": "Tax Summary",
                "blank_rows_above": 0,
                "columns": [
                    "Invoice Number","Invoice Date","Taxable Amount","CGST Input",
                    "SGST Input","IGST Input","CGST RCM","SGST RCM","IGST RCM","Invoice Amount",
                ],
            },
        ],
    },
    {
        "name": "ORIENTAL INSURANCE",
        "file": "ORIENTAL_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "Dept Description","Prod. Code Description","Pol.Type","Policy No.",
                    "Instl No.","Insured Name","Colln. No.","Cvr. Type",
                    "Premium / Ref. Premium (Rs.)","Comm. Rate","Comm. / Ref. Comm. (Rs.)",
                    "Policy Expiry Date","Bank Fin.","GST Exempt/Non-Exempt",
                ],
            },
        ],
    },
    {
        "name": "RELIANCE GENERAL",
        "file": "RELIANCE_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "Month","TransactionId","RegionName","BranchCode","BranchName","ProductCode",
                    "PolicyNumber","Endorsement","InsuredName","AgentCode","AgentName","AgentType",
                    "Discount","PremiumAmount","FinalIRDA","FinalIRDAComm","Reward","RewardCommission",
                    "TPPremiumAmount","FinalTP","FinalTPComm","TerrorismPremiumAmount","FinalTerrorism",
                    "FinalTerrorismComm","ChannelName","IRDAReqRefNo","RewardReqRefNo",
                    "TerrorismReqRefNo","TPReqRefNo","SkippedcovernoteY_N","POSPortalOSY_N",
                    "POLICY_COUNT","NoofLivesCovered","SUM_INSURED","NET_PREMIUM","CoverNoteNumber",
                    "CoinType","StateCode","StateName","Statement Month",
                ],
            },
        ],
    },
    {
        "name": "SHRIRAM GENERAL",
        "file": "SHRIRAM_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "DivisionName","SourceCode","SourceName","ProdDesc","Dpollno","Insuredname",
                    "GrossAmt","DOI","DOE","FinalDiscPer","TerrerismAmt","OD_TOTAL","TP_TOTAL",
                    "PremiumForComm","ODCommPer","ODCommAmt","TPCommPer","TPCommAmt","TotalCommAmt",
                    "CommMonth","RenewalType","VehicleNo","GSTNumber","State","VehfuelDesc","GVW",
                    "VehAge","PolicyIssunace Month(S_Date)",
                ],
            },
        ],
    },
    {
        "name": "TATA AIG",
        "file": "TATA_AIG_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Report",
                "blank_rows_above": 0,
                "columns": [
                    "Policy_No","Proposal_No","Business_Place","Policy_start_date","Policy_end_date",
                    "Policy_issue_date","Customer_name","Product_name","Package_name","Business_type",
                    "Record_type_desc","Record_type_cd","endorsement_no","OD_Premium",
                    "TP_Terrorism_premium","Total_Premium","OD_Commission","TP_Terrorism_commission",
                    "Total_commission","TDS","OD_Rate","TP_Rate","Approved_Grid_Per_Policy",
                    "sap_doc_posting_no","Partner_Application_No","Certificate_No","Sol_ID",
                    "Producer_Code","Subproducer_Code","Producer_Name","Commission_Cycle",
                    "UTR_Number","UTR_Date","Reason","Coinsurance_Flag","PANNO","b_code",
                    "report_accounting_period","CCOMMISSION_MONTH","table_name","GST_impact",
                    "Commission_Cycle_name","remarks","email","txt_commisiion_cycle","txt_status",
                    "statement_sent","Payment_Flag","num_category_id","txt_category","File_type",
                    "product_code","product_cd","Winflag","accounting_ts","sap_posting_ts",
                    "sap_doc_fiscal_year","statement_trig_ts",
                ],
            },
        ],
    },
    {
        "name": "UNITED INDIA",
        "file": "UNITED_INDIA_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "Sheet1",
                "blank_rows_above": 0,
                "columns": [
                    "Name","Agt/Brk/Cor_Ag_Code","Dep No","Policy No","Endorsement No",
                    "Insured Name","Policy Exp Dt","Premium Amt","Rate","Terrsm Loading","Comm Amt",
                ],
            },
        ],
    },
    {
        "name": "UNIVERSAL SOMPO",
        "file": "UNIVERSAL_SOMPO_Statement_FY2026-27.xlsx",
        "sheets": [
            {
                "name": "MOTOR",
                "blank_rows_above": 1,
                "columns": [
                    "ZONE","line_Of_Business","channel_Details","vertical_Channel_Map",
                    "policy_Issue_Month_Year","intermediary_Code","exception_applicable_grid",
                    "exception_id","sub_Imd_Name","intermediary_Category","txt_Office_Name",
                    "policyNo","endorsementNo","final_Policy_No","333.0","Vertical","intermediary",
                    "customer_Name","policy_Issue_Date","policy_start_date","policy_end_date",
                    "no_of_days","policy_issuance_state","breaking_flag","purchase_date",
                    "ilpos_Policy_Number","long_term_year","customer_Name","endorsement_Wording",
                    "registration_Date","age_Band","cubic_Capacity_Amp","vehicle_make","vehicle_model",
                    "yr_Of_Manufacturing","policyType","num_Product_Code","productName","product_Name",
                    "varient","gvw","policy_Status","tonnage","rtoLocation","rtoState",
                    "ncb_Applicability","customer_State","businessType","usgiNetPremium","netOdPremium",
                    "num_Net_Tp_Premium","consolidated_Amount_Multiyr","usgiNetPremium",
                    "terrorism_Premium","usgi_net_pr_less_tp_pr","net_Premium","saod_Satp","add_On_Flag",
                    "carrying_Capacity","seating_capacity","vehicle_model","eom_cluster","fuel_Type",
                    "age_Band","policy_duration","vehicle_age","motor_Registration_No",
                    "vehicle_Group_Name","state_code","rto_code","policy_type_new","state_Group_Name",
                    "declined_flag","health_fresh_premium","cover_Details","total_Si","co_Insurance_Type",
                    "txt_Ind_Corp_Flag","usgi_Branch_State","occupancy","preferred_referred",
                    "model_category","exception_id","grid_id","gwp","exception_gwp","final_gwp",
                    "gwp_slab","deviation_flag","product_flag","sub_channel_flag","grid_percentage",
                    "category_additional_percentage","exception_flat_percentage","irda_percentage",
                    "default_percentage","direct_code_flag","final_grid_percentage","computation",
                    "exception_additional_percentage","final_additional_percentage","loading_Percentage",
                    "discount_Percentage","additional_Percentage","rate","grid_amount_on_OD",
                    "grid_amount_on_TP","total_grid_amount","total_grid_amount_on_od_per",
                    "total_grid_amount_on_tp_per","net_total_grid_amount","grid_on_OD_pr_discount",
                    "grid_on_TP_pr_discount","grid_on_total_pr_discount","irda_od_perc","irda_tp_perc",
                    "irdai_amount_on_OD","irdai_amount_on_TP","total_irdai_amount",
                    "Instant Commission %","Instant Commission on OD","Instant Commission on TP",
                    "Total Instant Commission","final_grid_payable_on_OD","final_grid_payable_on_TP",
                    "final_grid_payable","total_amount","exception_applicable_grid","Final Remark",
                    "exception_additional_percentage","exception_applicable_on_od_net",
                    "special_condition_applicable","non_bolero_additional_percentage","special_deduction",
                    "non_bolero_contribution","ww","ww","Actual OD","Actual TP","Actual Total",
                    "OD","TP","TW Total","OD","TP","TW Total","OD","TP","TW Total",
                    "OD","TP","TW Total","OD","TP","TW Total",
                ],
            },
            {
                "name": "NON-MOTOR",
                "blank_rows_above": 1,
                "columns": [
                    "ZONE","line_Of_Business","channel_Details","vertical_Channel_Map",
                    "policy_Issue_Month_Year","intermediary_Code","exception_applicable_grid",
                    "exception_id","sub_Imd_Name","intermediary_Category","txt_Office_Name",
                    "policyNo","endorsementNo","final_Policy_No","intermediary","customer_Name",
                    "policy_Issue_Date","policy_start_date","policy_end_date","no_of_days",
                    "policy_issuance_state","breaking_flag","purchase_date","ilpos_Policy_Number",
                    "long_term_year","customer_Name","endorsement_Wording","registration_Date",
                    "age_Band","cubic_Capacity_Amp","vehicle_make","vehicle_model","yr_Of_Manufacturing",
                    "policyType","num_Product_Code","productName","product_Name","varient","gvw",
                    "policy_Status","tonnage","rtoLocation","rtoState","ncb_Applicability",
                    "customer_State","businessType","usgiNetPremium","netOdPremium","num_Net_Tp_Premium",
                    "consolidated_Amount_Multiyr","usgiNetPremium","terrorism_Premium","TERRERISM_SHARE",
                    "usgi_net_pr_less_tp_pr","net_Premium","saod_Satp","add_On_Flag","carrying_Capacity",
                    "seating_capacity","vehicle_model","eom_cluster","fuel_Type","age_Band",
                    "policy_duration","vehicle_age","motor_Registration_No","vehicle_Group_Name",
                    "state_code","rto_code","policy_type_new","state_Group_Name","declined_flag",
                    "health_fresh_premium","cover_Details","total_Si","co_Insurance_Type",
                    "txt_Ind_Corp_Flag","usgi_Branch_State","occupancy","preferred_referred",
                    "model_category","exception_id","grid_id","gwp","slab","exception_gwp","final_gwp",
                    "gwp_slab","deviation_flag","product_flag","sub_channel_flag","grid_percentage",
                    "category_additional_percentage","exception_flat_percentage","irda_percentage",
                    "default_percentage","direct_code_flag","Manual final_grid_percentage","computation",
                    "exception_additional_percentage","final_additional_percentage","loading_Percentage",
                    "discount_Percentage","additional_Percentage","rate","grid_amount_on_OD",
                    "grid_amount_on_TP","total_grid_amount","total_grid_amount_on_od_per",
                    "total_grid_amount_on_tp_per","net_total_grid_amount","grid_on_OD_pr_discount",
                    "grid_on_TP_pr_discount","grid_on_total_pr_discount","irda_od_perc","irda_tp_perc",
                    "irdai_amount_on_OD","irdai_amount_on_TP","total_irdai_amount","5 % TERRERISM_SHARE",
                    "final_grid_payable_on_OD","final_grid_payable_on_TP","final_grid_payable",
                    "total_amount","exception_applicable_grid","exception_additional_percentage",
                    "exception_applicable_on_od_net","special_condition_applicable",
                    "non_bolero_additional_percentage","special_deduction","non_bolero_contribution",
                    "remarks","concat_column","swing_grid","Channel Category","TOTAL_SERVICE_TAX",
                    "GST REMARK FOR INVOICE","remarks",
                ],
            },
        ],
    },
]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    summary = []

    for cfg in INSURERS:
        insurer   = cfg["name"]
        filename  = cfg["file"]
        sheets    = cfg["sheets"]

        print(f"\nBuilding: {insurer} ...")
        try:
            wb = Workbook()
            # Remove default empty sheet
            wb.remove(wb.active)

            for sheet_cfg in sheets:
                write_sheet(
                    wb,
                    sheet_cfg["name"],
                    sheet_cfg["columns"],
                    sheet_cfg["blank_rows_above"],
                )

            out_path = os.path.join(OUT_DIR, filename)
            wb.save(out_path)

            sheet_names = "/".join(s["name"] for s in sheets)
            col_counts  = "/".join(str(len(s["columns"])) for s in sheets)
            header_rows = "/".join(str(s["blank_rows_above"] + 1) for s in sheets)
            status      = "OK"
            print(f"  -> {filename}  sheets={sheet_names}  cols={col_counts}")

        except Exception:
            sheet_names = "—"
            col_counts  = "—"
            header_rows = "—"
            status      = "FAILED"
            print(f"  [ERROR] {insurer}:")
            traceback.print_exc()

        summary.append((insurer, filename, sheet_names, header_rows, col_counts, status))

    # Summary table
    print("\n")
    SEP = "=" * 118
    print(SEP)
    print(f"{'Insurer':<22} {'File':<40} {'Sheet(s)':<22} {'Hdr Row':<9} {'Cols':<8} Status")
    print("-" * 118)
    for insurer, fname, sheets, hrow, cols, status in summary:
        flag = "  <-- FAILED" if status == "FAILED" else ""
        print(f"{insurer:<22} {fname:<40} {sheets:<22} {hrow:<9} {cols:<8} {status}{flag}")
    print(SEP)

    ok     = sum(1 for r in summary if r[5] == "OK")
    failed = sum(1 for r in summary if r[5] == "FAILED")
    print(f"\nResult: {ok} succeeded, {failed} failed.\n")
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
