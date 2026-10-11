"""
Split CaliberPO_Consolidated_Grid_Extractv2.xlsx into one .xlsx per insurer,
upload each to Grid Import, and call Main ?action=importGrids.
"""
from __future__ import annotations

import json
import os
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import requests
from openpyxl import Workbook

SRC = Path(r"c:\Users\Rajendra\Downloads\CaliberPO_Consolidated_Grid_Extractv2.xlsx")
OUT_DIR = Path(r"d:\HRSrini\caliber-po\masters\grid_import_split")
GRID_IMPORT = "1-h1bFSfpP8yZxejdK95IEXXbmfjnbBb7"
MAIN = "https://script.google.com/macros/s/AKfycbyTM-OpjCtov-a_CCRaN9qZZ9n2xWkHd_w0TDKUtAdk3Ez8bkCDymsRaVldPMcVb9XK/exec"
KEY = "caliberpo-aps-bootstrap-2026"
PRODUCT_TABS = ("TW_26-27", "PCV_26-27", "GCV_26-27", "MiscD_26-27", "PC_26-27")


def oauth_token() -> str:
    clasprc = Path(os.environ["USERPROFILE"]) / ".clasprc.json"
    rc = json.loads(clasprc.read_text(encoding="utf-8"))
    t = rc["tokens"]["default"]
    r = requests.post(
        "https://oauth2.googleapis.com/token",
        data={
            "client_id": t["client_id"],
            "client_secret": t["client_secret"],
            "refresh_token": t["refresh_token"],
            "grant_type": "refresh_token",
        },
        timeout=60,
    )
    r.raise_for_status()
    return r.json()["access_token"]


def analyze(path: Path) -> None:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        if name not in PRODUCT_TABS and not name.endswith("_26-27"):
            if name in ("Still Missing",) or name.startswith("READ"):
                ws = wb[name]
                print(f"--- {name} ---")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 20:
                        break
                    print(row)
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = list(rows[0])
        ic = hdr.index("insurer_code")
        company = hdr.index("insurance_company")
        codes = Counter()
        for r in rows[1:]:
            if not r or r[ic] is None or str(r[ic]).strip() == "":
                continue
            codes[str(r[ic]).strip()] += 1
        print(f"{name}: rows={sum(codes.values())} insurers={dict(codes)}")
    wb.close()


def split_by_insurer(path: Path, out_dir: Path) -> dict[str, Path]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # code -> tab -> (header, rows)
    by_code: dict[str, dict[str, tuple[list, list]]] = defaultdict(dict)
    names: dict[str, str] = {}

    for tab in wb.sheetnames:
        if not any(tab.startswith(p.split("_")[0]) for p in PRODUCT_TABS) and not tab.endswith("_26-27"):
            continue
        if tab.startswith("READ") or tab == "Still Missing":
            continue
        ws = wb[tab]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        hdr = list(rows[0])
        try:
            ic = hdr.index("insurer_code")
            company = hdr.index("insurance_company")
        except ValueError:
            continue
        buckets: dict[str, list] = defaultdict(list)
        for r in rows[1:]:
            if not r or r[ic] is None or str(r[ic]).strip() == "":
                continue
            code = str(r[ic]).strip()
            buckets[code].append([("" if v is None else v) for v in r])
            if code not in names and r[company]:
                names[code] = str(r[company]).strip()
        for code, data in buckets.items():
            by_code[code][tab] = (hdr, data)

    wb.close()
    out_dir.mkdir(parents=True, exist_ok=True)
    written: dict[str, Path] = {}
    for code, tabs in sorted(by_code.items()):
        out = out_dir / f"{code}_grid_26-27.xlsx"
        nwb = Workbook()
        # remove default
        nwb.remove(nwb.active)
        total = 0
        for tab, (hdr, data) in sorted(tabs.items()):
            ws = nwb.create_sheet(tab)
            ws.append(hdr)
            for row in data:
                # trim/pad to header width
                row = list(row)[: len(hdr)]
                while len(row) < len(hdr):
                    row.append("")
                ws.append(row)
            total += len(data)
        nwb.save(out)
        written[code] = out
        print(f"wrote {out.name}: tabs={list(tabs)} rows={total} name={names.get(code,'')}")
    return written


def upload(token: str, path: Path, parent: str) -> str:
    meta = {"name": path.name, "parents": [parent]}
    r = requests.post(
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id,name",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "metadata": ("metadata", json.dumps(meta), "application/json"),
            "file": (
                path.name,
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        timeout=300,
    )
    print("upload", path.name, r.status_code, r.text[:200])
    r.raise_for_status()
    return r.json()["id"]


def main() -> int:
    if not SRC.exists():
        print("missing source", SRC)
        return 1
    print("=== analyze ===")
    analyze(SRC)
    print("=== split ===")
    files = split_by_insurer(SRC, OUT_DIR)
    if not files:
        print("no insurer files produced")
        return 1
    print("=== upload ===")
    token = oauth_token()
    for code, path in files.items():
        upload(token, path, GRID_IMPORT)
    print("=== importGrids ===")
    # Prefer dedicated action; fall back message if missing
    r = requests.get(
        MAIN,
        params={"action": "importGrids", "key": KEY},
        timeout=540,
        headers={"User-Agent": "caliberpo"},
    )
    print("status", r.status_code)
    print(r.text[:2000])
    return 0


if __name__ == "__main__":
    sys.exit(main())
